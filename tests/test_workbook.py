"""Workbook format tests (CLAUDE.md section 14, "Workbook format").

These run against the real reference export wherever possible. A synthetic
file that has been tidied up would prove nothing about the format the
instrument actually writes.
"""

from __future__ import annotations

from datetime import datetime

import pytest
from openpyxl import load_workbook

from maxone_loop.metrics.workbook import MISSING, MetricsWorkbook, WorkbookError
from maxone_loop.reason_codes import COLUMN_MISSING, NUMERIC_COERCION_FAILED, SHEET_SET_MISMATCH
from maxone_loop.synthetic import SyntheticRecording, write_workbook

EXPECTED_WELLPLATES = {
    "P005163", "P005157", "P005124", "P005232", "P005211", "P005169", "P005190",
}


def test_reference_workbook_yields_seven_recordings(reference_workbook, runnable_config):
    schema = runnable_config().metrics_schema
    workbook = MetricsWorkbook(reference_workbook, schema)
    recordings = workbook.recordings()

    assert len(recordings) == 7
    assert {r.identity.wellplate_id for r in recordings.values()} == EXPECTED_WELLPLATES


def test_each_recording_has_exactly_one_activity_instance(reference_workbook, runnable_config):
    schema = runnable_config().metrics_schema
    workbook = MetricsWorkbook(reference_workbook, schema)

    for recording in workbook.recordings().values():
        assert len(recording.instances["Activity Analysis"]) == 1
        assert len(recording.instances["ISI-N Burst Detector"]) == 1


def test_analysis_type_resolved_without_instance_parity(reference_workbook, runnable_config):
    """Parity happens to hold in the reference file; the reader must not use it."""
    schema = runnable_config().metrics_schema
    workbook = MetricsWorkbook(reference_workbook, schema)
    types = workbook._analysis_types()

    # The coincidence exists...
    assert all(
        (instance % 2 == 1) == (kind == "Activity Analysis") for instance, kind in types.items()
    )
    # ...and the mapping the reader actually uses comes from Analysis Parameters.
    assert types[1] == "Activity Analysis"
    assert types[2] == "ISI-N Burst Detector"


def test_na_is_missing_never_zero(reference_workbook, runnable_config):
    """P005211 has "N/A" dispersion. It must not read as 0.0."""
    schema = runnable_config().metrics_schema
    workbook = MetricsWorkbook(reference_workbook, schema)
    row = workbook.activity_row(9)

    assert row["Firing Rate std [Hz]"] is MISSING
    assert row["Firing Rate CV"] is MISSING
    assert workbook.as_float(row["Firing Rate std [Hz]"], where="test") is None
    # The mean itself is present and numeric.
    assert workbook.as_float(row["Mean Firing Rate [Hz]"], where="test") == 0.19


def test_unnamed_index_column_is_dropped_not_shifted(reference_workbook, runnable_config):
    """Network - Burst Level column A is a DataFrame index remnant."""
    schema = runnable_config().metrics_schema
    workbook = MetricsWorkbook(reference_workbook, schema)
    sheet = workbook.sheet("network_burst_level")

    raw = load_workbook(reference_workbook, read_only=True, data_only=True)
    try:
        raw_header = list(next(raw["Network - Burst Level"].iter_rows(values_only=True)))
    finally:
        raw.close()

    assert raw_header[0] is None  # the remnant
    assert sheet.headers[0] == "Instance"  # and the reader starts at the real first column
    assert None not in sheet.headers
    assert len(sheet.headers) == len(raw_header) - 1


def test_burst_sheet_is_not_a_metadata_source(reference_workbook, runnable_config):
    """Metadata is split across rows there; Meta Data is the only source."""
    schema = runnable_config().metrics_schema
    assert schema.known_quirks.network_burst_level.use_as_metadata_source is False

    workbook = MetricsWorkbook(reference_workbook, schema)
    rows = workbook.sheet("network_burst_level").rows

    # Row 0 of the first block carries the identity but not the group columns.
    assert rows[0]["Wellplate ID"] == "P005163"
    assert rows[0]["Well Group Name"] is MISSING
    # Row 1 carries the group columns but not the identity.
    assert rows[1]["Well Group Name"] == "Default Group"
    assert rows[1]["Wellplate ID"] is MISSING

    # Meta Data has every field on every row.
    for recording in workbook.recordings().values():
        assert recording.identity.wellplate_id
        assert recording.identity.assay_run_id


def test_zero_burst_recording_does_not_crash(reference_workbook, runnable_config):
    """P005163 has no bursts but still emits placeholder rows."""
    schema = runnable_config().metrics_schema
    workbook = MetricsWorkbook(reference_workbook, schema)

    network = workbook.sheet("network_well_level")
    burst_row = next(r for r in network.rows if r["Instance"] == 2)
    assert burst_row["Burst Frequency [Hz]"] == 0
    assert burst_row["Mean Spikes per Burst"] is MISSING

    # Two placeholder rows exist, so the row count is not a burst count.
    placeholders = [r for r in workbook.sheet("network_burst_level").rows if r["burstTime"] is MISSING]
    assert len(placeholders) >= 2


def test_dates_and_times_are_strings_not_serials(reference_workbook, runnable_config):
    schema = runnable_config().metrics_schema
    workbook = MetricsWorkbook(reference_workbook, schema)
    meta = workbook.sheet("meta_data").rows[0]

    assert meta["Date"] == "2026-04-09"
    assert meta["Start Time"] == "13:56:25"

    recording = workbook.recordings()[meta["Folder Path"]]
    assert recording.start == datetime(2026, 4, 9, 13, 56, 25)
    assert recording.stop == datetime(2026, 4, 9, 14, 1, 25)


def test_durations_are_not_exactly_300(reference_workbook, runnable_config):
    schema = runnable_config().metrics_schema
    workbook = MetricsWorkbook(reference_workbook, schema)
    durations = {r.acquisition.duration_seconds for r in workbook.recordings().values()}

    assert all(d != 300.0 for d in durations)
    assert min(durations) >= 300.0 and max(durations) <= 301.0


@pytest.mark.parametrize("sheet", ["Meta Data", "Analysis Parameters", "Activity - Well Level"])
def test_missing_required_sheet_is_rejected(tmp_path, runnable_config, sheet):
    schema = runnable_config().metrics_schema
    path = write_workbook(
        tmp_path / f"metrics_data_20260101_0000{abs(hash(sheet)) % 100:02d}.xlsx",
        [SyntheticRecording("P1", "000001", 0.5)],
    )
    book = load_workbook(path)
    del book[sheet]
    book.save(path)

    with pytest.raises(WorkbookError) as exc:
        MetricsWorkbook(path, schema)
    assert exc.value.reason_code == SHEET_SET_MISMATCH
    assert "required" in exc.value.message


def test_summary_metrics_export_is_accepted(tmp_path, runnable_config):
    """"Export summary metrics" omits the burst-level sheet.

    The decision path never reads it, so a summary export must still work --
    requiring it would make a correct export unreadable. See CLAUDE.md section 3.
    """
    config = runnable_config(
        **{
            "metrics_schema.recording_selection.strategy": "folder_path_prefix",
            "metrics_schema.recording_selection.folder_path": "/home/mxwbio/Data/Synthetic/",
        }
    )
    path = write_workbook(
        tmp_path / "metrics_data_20260101_000010.xlsx",
        [SyntheticRecording("P1", "000001", 0.5)],
    )
    book = load_workbook(path)
    del book["Network - Burst Level"]
    book.save(path)

    workbook = MetricsWorkbook(path, config.metrics_schema)
    assert workbook.has_sheet("network_burst_level") is False
    assert workbook.has_sheet("activity_well_level") is True
    assert len(workbook.recordings()) == 1

    # And the metric is still readable end to end.
    from maxone_loop.metrics.extract import extract
    from maxone_loop.metrics.selection import select_recording

    selection = select_recording(workbook, config.metrics_schema)
    result = extract(workbook, config.metrics_schema, selection, workbook_sha256="0" * 64)
    assert result.mean_firing_frequency_hz == 0.5


def test_activity_only_export_is_accepted(tmp_path, runnable_config):
    """A Network Assay analysed with Activity Analysis alone has no Network sheets."""
    config = runnable_config(
        **{
            "metrics_schema.recording_selection.strategy": "folder_path_prefix",
            "metrics_schema.recording_selection.folder_path": "/home/mxwbio/Data/Synthetic/",
        }
    )
    path = write_workbook(
        tmp_path / "metrics_data_20260101_000011.xlsx",
        [SyntheticRecording("P1", "000001", 0.5)],
    )
    book = load_workbook(path)
    del book["Network - Burst Level"]
    del book["Network - Well Level"]
    book.save(path)

    workbook = MetricsWorkbook(path, config.metrics_schema)
    assert workbook.present_sheets == (
        "activity_well_level",
        "analysis_parameters",
        "meta_data",
    )
    assert len(workbook.recordings()) == 1


def test_reading_an_absent_optional_sheet_raises_rather_than_returning_empty(
    tmp_path, runnable_config
):
    schema = runnable_config().metrics_schema
    path = write_workbook(
        tmp_path / "metrics_data_20260101_000012.xlsx",
        [SyntheticRecording("P1", "000001", 0.5)],
    )
    book = load_workbook(path)
    del book["Network - Burst Level"]
    book.save(path)

    workbook = MetricsWorkbook(path, schema)
    with pytest.raises(WorkbookError) as exc:
        workbook.sheet("network_burst_level")
    assert exc.value.reason_code == SHEET_SET_MISMATCH


def test_unexpected_extra_sheet_is_rejected(tmp_path, runnable_config):
    schema = runnable_config().metrics_schema
    path = write_workbook(
        tmp_path / "metrics_data_20260101_000001.xlsx",
        [SyntheticRecording("P1", "000001", 0.5)],
    )
    book = load_workbook(path)
    book.create_sheet("Unexpected")
    book.save(path)

    with pytest.raises(WorkbookError) as exc:
        MetricsWorkbook(path, schema)
    assert exc.value.reason_code == SHEET_SET_MISMATCH


def test_renamed_metric_column_is_rejected(tmp_path, runnable_config):
    config = runnable_config()
    path = write_workbook(
        tmp_path / "metrics_data_20260101_000002.xlsx",
        [SyntheticRecording("P1", "000001", 0.5)],
    )
    book = load_workbook(path)
    sheet = book["Activity - Well Level"]
    for cell in sheet[1]:
        if cell.value == "Mean Firing Rate [Hz]":
            # A unit change is exactly this: same concept, different meaning.
            cell.value = "Mean Firing Rate [spikes/min]"
    book.save(path)

    workbook = MetricsWorkbook(path, config.metrics_schema)
    with pytest.raises(WorkbookError) as exc:
        workbook.sheet("activity_well_level").require(
            config.metrics_schema.decision_metric.column
        )
    assert exc.value.reason_code == COLUMN_MISSING


def test_non_numeric_rate_is_a_coercion_failure(runnable_config):
    schema = runnable_config().metrics_schema
    assert schema.file_format.numeric_coercion == "reject"

    with pytest.raises(WorkbookError) as exc:
        MetricsWorkbook.as_float("not a number", where="test")
    assert exc.value.reason_code == NUMERIC_COERCION_FAILED


@pytest.mark.parametrize("value", [float("nan"), float("inf"), float("-inf")])
def test_nan_and_infinity_are_rejected(value):
    with pytest.raises(WorkbookError) as exc:
        MetricsWorkbook.as_float(value, where="test")
    assert exc.value.reason_code == NUMERIC_COERCION_FAILED


def test_truncated_file_is_reported_not_crashed(tmp_path, runnable_config):
    schema = runnable_config().metrics_schema
    path = tmp_path / "metrics_data_20260101_000003.xlsx"
    path.write_bytes(b"PK\x03\x04 truncated")

    with pytest.raises(WorkbookError):
        MetricsWorkbook(path, schema)
