"""CLI tests, including the exact invocations documented in CLAUDE.md section 14."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
import yaml

from maxone_loop.cli import main

REPO_ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def config_dir(tmp_path, runnable_config):
    """Write a runnable configuration to disk and return its directory."""

    def build(**extra):
        config = runnable_config(**extra)
        directory = tmp_path / "cfg"
        directory.mkdir(exist_ok=True)
        for name, section in (
            ("experiment.yaml", config.experiment),
            ("metrics_schema.yaml", config.metrics_schema),
            ("stimulation_protocols.yaml", config.stimulation_protocols),
        ):
            (directory / name).write_text(
                yaml.safe_dump(section.model_dump(mode="json"), sort_keys=False)
            )
        return directory

    return build


def test_validate_config_on_the_tracked_files(capsys):
    assert main(["validate-config", "--config", str(REPO_ROOT / "config" / "experiment.yaml")]) == 0
    out = capsys.readouterr().out
    assert "configuration OK" in out
    assert "armed mode blocked" in out


def test_config_flag_works_before_the_subcommand(capsys):
    assert main(["--config", str(REPO_ROOT / "config"), "validate-config"]) == 0
    assert "configuration OK" in capsys.readouterr().out


def test_inspect_lists_every_recording(capsys, reference_workbook):
    code = main(
        [
            "inspect",
            "--config",
            str(REPO_ROOT / "config" / "experiment.yaml"),
            str(reference_workbook),
        ]
    )
    out = capsys.readouterr().out

    assert code == 0
    assert "7 recording(s)" in out
    for wellplate in ("P005163", "P005157", "P005211", "P005190"):
        assert wellplate in out
    assert "export time" in out


def test_replay_with_the_tracked_config_refuses_to_guess(capsys):
    """strategy is null there, so replay must fail closed rather than pick a row."""
    code = main(
        [
            "replay",
            "--config",
            str(REPO_ROOT / "config" / "experiment.yaml"),
            "--input",
            str(Path(__file__).parent / "fixtures"),
        ]
    )
    out = capsys.readouterr().out

    assert code == 1
    assert "RECORDING_NOT_FOUND" in out
    assert "refusing to guess" in out


def test_replay_selects_and_passes_qc(capsys, config_dir, tmp_path, reference_workbook):
    directory = config_dir()
    fixtures = tmp_path / "fixtures"
    fixtures.mkdir()
    shutil.copy(reference_workbook, fixtures / reference_workbook.name)

    code = main(["replay", "--config", str(directory), "--input", str(fixtures)])
    out = capsys.readouterr().out

    assert code == 0
    assert "P005157" in out
    assert "-> PASS" in out


def test_replay_rejects_the_single_electrode_recording(capsys, config_dir, tmp_path, reference_workbook):
    directory = config_dir(
        **{
            "metrics_schema.recording_selection.wellplate_id": "P005211",
            "metrics_schema.recording_selection.assay_run_id": "000031",
        }
    )
    fixtures = tmp_path / "fixtures"
    fixtures.mkdir()
    shutil.copy(reference_workbook, fixtures / reference_workbook.name)

    code = main(["replay", "--config", str(directory), "--input", str(fixtures)])
    out = capsys.readouterr().out

    assert code == 1
    assert "REJECT" in out
    assert "SINGLE_ELECTRODE_MEAN" in out


def test_replay_refuses_armed_mode(capsys, config_dir, tmp_path):
    directory = config_dir(**{"experiment.experiment.mode": "armed"})
    assert main(["replay", "--config", str(directory), "--input", str(tmp_path)]) == 2
    assert "armed" in capsys.readouterr().err


def test_run_dry_run_completes(capsys, config_dir, tmp_path):
    from maxone_loop.synthetic import SyntheticRecording, write_workbook

    directory = config_dir(
        **{
            "metrics_schema.recording_selection.strategy": "folder_path_prefix",
            "metrics_schema.recording_selection.folder_path": "/home/mxwbio/Data/Synthetic/",
        }
    )
    inbox = tmp_path / "inbox"
    write_workbook(
        inbox / "metrics_data_20260101_000000.xlsx",
        [SyntheticRecording("P1", "000001", 2.0)],
    )

    code = main(
        ["run", "--config", str(directory), "--mode", "dry_run", "--input", str(inbox)]
    )
    out = capsys.readouterr().out

    assert code == 0
    assert "final state : COMPLETE" in out
    assert "THRESHOLD_REACHED" in out
    assert "RECORD" in out  # the baseline cycle


def test_run_without_a_watch_directory_is_an_error(capsys):
    code = main(
        ["run", "--config", str(REPO_ROOT / "config" / "experiment.yaml"), "--mode", "dry_run"]
    )
    assert code == 2
    assert "metrics_watch_directory is null" in capsys.readouterr().err


def test_replay_descends_into_export_subfolders(capsys, config_dir, tmp_path, reference_workbook):
    """Exports land in a timestamped subfolder, so a flat glob would miss them."""
    directory = config_dir()
    root = tmp_path / "exports"
    export = root / "20260409_150425"
    export.mkdir(parents=True)
    shutil.copy(reference_workbook, export / reference_workbook.name)

    code = main(["replay", "--config", str(directory), "--input", str(root)])
    out = capsys.readouterr().out

    assert code == 0
    assert "P005157" in out


def test_inspect_reports_absent_optional_sheets(capsys, config_dir, tmp_path, reference_workbook):
    from openpyxl import load_workbook

    directory = config_dir()
    path = tmp_path / reference_workbook.name
    shutil.copy(reference_workbook, path)
    book = load_workbook(path)
    del book["Network - Burst Level"]
    book.save(path)

    code = main(["inspect", "--config", str(directory), str(path)])
    out = capsys.readouterr().out

    assert code == 0
    assert "optional sheet(s) not in this export" in out
    assert "network_burst_level" in out
    assert "7 recording(s)" in out


def test_inspect_warns_when_activity_analysis_is_missing(capsys, config_dir, tmp_path):
    """Running only Network Analysis leaves no metric to read."""
    from openpyxl import load_workbook

    from maxone_loop.synthetic import SyntheticRecording, write_workbook

    directory = config_dir()
    path = write_workbook(
        tmp_path / "metrics_data_20260101_000000.xlsx",
        [SyntheticRecording("P1", "000001", 0.5)],
    )
    book = load_workbook(path)
    for row in book["Analysis Parameters"].iter_rows(min_row=2):
        if row[1].value == "Activity Analysis":
            row[1].value = "ISI-N Burst Detector"
    book.save(path)

    code = main(["inspect", "--config", str(directory), str(path)])
    out = capsys.readouterr().out

    assert code == 0
    assert "WARNING" in out
    assert "Activity Analysis" in out
