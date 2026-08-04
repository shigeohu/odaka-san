"""Read a MaxLab Live metrics workbook.

The reader is deliberately pedantic. Every rule here exists because the real
export violates a reasonable assumption somewhere -- see CLAUDE.md section 2.5.

Nothing in this module writes to the workbook.
"""

from __future__ import annotations

import math
from dataclasses import dataclass, field
from datetime import date as date_cls
from datetime import datetime, time
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook

from ..config.models import MetricsSchema
from ..reason_codes import (
    COLUMN_MISSING,
    NUMERIC_COERCION_FAILED,
    SHEET_SET_MISMATCH,
    WORKBOOK_UNREADABLE,
)

#: Sentinel for a cell whose content is one of the configured missing markers.
MISSING = object()


class WorkbookError(Exception):
    """The workbook could not be read under the configured schema."""

    def __init__(self, reason_code: str, message: str) -> None:
        super().__init__(message)
        self.reason_code = reason_code
        self.message = message


@dataclass(frozen=True)
class RecordingIdentity:
    """Identifies one recording inside a workbook.

    ``folder_path`` is the primary key. The structured fields are carried for
    provenance and for the ``wellplate_and_run_id`` selection strategy;
    ``assay_run_id`` alone is not unique across wellplates.
    """

    folder_path: str
    wellplate_id: str
    assay_run_id: str
    well_number: int

    def as_dict(self) -> dict[str, Any]:
        return {
            "folder_path": self.folder_path,
            "wellplate_id": self.wellplate_id,
            "assay_run_id": self.assay_run_id,
            "well_number": self.well_number,
        }


@dataclass(frozen=True)
class Acquisition:
    duration_seconds: float | None
    number_of_configurations: int | None
    sampling_frequency_hz: float | None
    gain: float | None
    lsb_uv: float | None
    hpf_hz: float | None


@dataclass(frozen=True)
class Recording:
    """One recording and the analysis instances derived from it."""

    identity: RecordingIdentity
    acquisition: Acquisition
    start: datetime | None
    stop: datetime | None
    #: analysis type -> Instance ids. Normally one id per type, but the reader
    #: reports what it finds; selection decides whether that is acceptable.
    instances: dict[str, tuple[int, ...]] = field(default_factory=dict)


@dataclass(frozen=True)
class Sheet:
    """A sheet reduced to header-keyed rows.

    Rows are dicts keyed by header name. ``Network - Burst Level`` carries an
    unnamed index column in position A; its ``None`` header is dropped here so
    that no caller can read it by accident.
    """

    name: str
    headers: tuple[str, ...]
    rows: tuple[dict[str, Any], ...]

    def require(self, *columns: str) -> None:
        absent = [c for c in columns if c not in self.headers]
        if absent:
            raise WorkbookError(
                COLUMN_MISSING,
                f"sheet {self.name!r} is missing column(s): {', '.join(map(repr, absent))}",
            )


class MetricsWorkbook:
    """A validated, read-only view of one metrics export."""

    def __init__(self, path: Path, schema: MetricsSchema) -> None:
        self.path = Path(path)
        self.schema = schema
        self._sheets: dict[str, Sheet] = {}
        self._load()

    # -- loading ------------------------------------------------------------

    def _load(self) -> None:
        try:
            workbook = load_workbook(self.path, read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises a wide range of types
            raise WorkbookError(
                WORKBOOK_UNREADABLE, f"cannot open {self.path.name}: {exc}"
            ) from exc

        try:
            sheets = self.schema.sheets
            actual = set(workbook.sheetnames)
            required = sheets.required_names()
            optional = sheets.optional_names()

            missing = set(required.values()) - actual
            if missing:
                raise WorkbookError(
                    SHEET_SET_MISMATCH,
                    f"{self.path.name}: missing required sheet(s) {sorted(missing)}",
                )

            # An unknown sheet means the export format changed under us.
            if sheets.reject_unknown_sheets:
                unknown = actual - set(required.values()) - set(optional.values())
                if unknown:
                    raise WorkbookError(
                        SHEET_SET_MISMATCH,
                        f"{self.path.name}: unexpected sheet(s) {sorted(unknown)}",
                    )

            # An absent optional sheet is a legitimate export choice, not a
            # defect: "Export summary metrics" omits the burst-level sheet, and
            # the Network sheets exist only when a Network Analysis trial was
            # included. See CLAUDE.md section 3.
            for key, name in {**required, **optional}.items():
                if name in actual:
                    self._sheets[key] = self._read_sheet(workbook[name], name)
        finally:
            workbook.close()

    def _read_sheet(self, worksheet: Any, name: str) -> Sheet:
        rows = worksheet.iter_rows(values_only=True)
        header_row = self.schema.file_format.header_row
        raw_header: tuple[Any, ...] | None = None
        for index, row in enumerate(rows):
            if index == header_row:
                raw_header = row
                break
        if raw_header is None:
            raise WorkbookError(WORKBOOK_UNREADABLE, f"sheet {name!r} has no header row")

        # An unnamed header is the DataFrame index remnant in Network - Burst
        # Level. Drop the column entirely rather than inventing a name for it.
        keep = [(i, str(h)) for i, h in enumerate(raw_header) if h is not None and str(h) != ""]
        headers = tuple(h for _, h in keep)
        if len(set(headers)) != len(headers):
            duplicates = sorted({h for h in headers if headers.count(h) > 1})
            raise WorkbookError(
                COLUMN_MISSING,
                f"sheet {name!r} has duplicate header(s) {duplicates}; the column "
                "mapping would be ambiguous",
            )

        parsed = tuple(
            {header: self._normalise(row[i] if i < len(row) else None) for i, header in keep}
            for row in rows
        )
        return Sheet(name=name, headers=headers, rows=parsed)

    def _normalise(self, value: Any) -> Any:
        """Map the configured missing markers onto ``MISSING``.

        ``"N/A"`` is a string sitting inside otherwise-numeric columns. It must
        never silently become 0.
        """
        if value is None:
            return MISSING
        if isinstance(value, str):
            stripped = value.strip()
            if stripped in self.schema.file_format.missing_markers:
                return MISSING
            return stripped
        return value

    # -- access -------------------------------------------------------------

    def sheet(self, key: str) -> Sheet:
        """The sheet for ``key``. Raises if it was not present in the export."""
        try:
            return self._sheets[key]
        except KeyError:
            raise WorkbookError(
                SHEET_SET_MISMATCH,
                f"{self.path.name}: sheet {key!r} is not present in this export",
            ) from None

    def has_sheet(self, key: str) -> bool:
        return key in self._sheets

    @property
    def present_sheets(self) -> tuple[str, ...]:
        return tuple(sorted(self._sheets))

    # -- typed cell readers -------------------------------------------------

    @staticmethod
    def as_float(value: Any, *, where: str) -> float | None:
        """Return a finite float, ``None`` for missing, or raise.

        A value that looks numeric but is not is a coercion failure, never a
        dropped row (``numeric_coercion: reject``).
        """
        if value is MISSING:
            return None
        if isinstance(value, bool):
            raise WorkbookError(NUMERIC_COERCION_FAILED, f"{where}: boolean is not a number")
        if isinstance(value, (int, float)):
            number = float(value)
        elif isinstance(value, str):
            try:
                number = float(value)
            except ValueError as exc:
                raise WorkbookError(
                    NUMERIC_COERCION_FAILED, f"{where}: {value!r} is not numeric"
                ) from exc
        else:
            raise WorkbookError(
                NUMERIC_COERCION_FAILED, f"{where}: {type(value).__name__} is not numeric"
            )
        if math.isnan(number) or math.isinf(number):
            raise WorkbookError(NUMERIC_COERCION_FAILED, f"{where}: {number} is not finite")
        return number

    @staticmethod
    def as_int(value: Any, *, where: str) -> int | None:
        number = MetricsWorkbook.as_float(value, where=where)
        if number is None:
            return None
        if number != int(number):
            raise WorkbookError(NUMERIC_COERCION_FAILED, f"{where}: {number} is not an integer")
        return int(number)

    @staticmethod
    def as_str(value: Any, *, where: str) -> str | None:
        if value is MISSING:
            return None
        if isinstance(value, str):
            return value
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            # Assay Run ID is "000030" in the reference file, but a workbook
            # written by another tool could store it as a number.
            return str(int(value)) if float(value).is_integer() else str(value)
        raise WorkbookError(NUMERIC_COERCION_FAILED, f"{where}: cannot read {value!r} as text")

    def as_datetime(self, date_value: Any, time_value: Any, *, where: str) -> datetime | None:
        """Combine the string date and time columns.

        MaxLab Live writes these as strings, not Excel serials, but a workbook
        round-tripped through Excel may carry real date/time objects. Both are
        accepted; anything else fails closed.
        """
        if date_value is MISSING or time_value is MISSING:
            return None

        fmt = self.schema.file_format
        try:
            if isinstance(date_value, datetime):
                day = date_value.date()
            elif isinstance(date_value, date_cls):
                day = date_value
            else:
                day = datetime.strptime(str(date_value), fmt.date_format).date()

            if isinstance(time_value, datetime):
                clock = time_value.time()
            elif isinstance(time_value, time):
                clock = time_value
            else:
                clock = datetime.strptime(str(time_value), fmt.time_format).time()
        except ValueError as exc:
            raise WorkbookError(
                NUMERIC_COERCION_FAILED,
                f"{where}: cannot parse {date_value!r} {time_value!r} as a timestamp",
            ) from exc
        return datetime.combine(day, clock)

    # -- recordings ---------------------------------------------------------

    def recordings(self) -> dict[str, Recording]:
        """Build the recording table from ``Meta Data``.

        One recording produces several analysis instances; they are grouped by
        ``Folder Path`` and their analysis types resolved through
        ``Analysis Parameters``. Instance parity is never consulted.
        """
        schema = self.schema
        meta = self.sheet("meta_data")
        ident = schema.recording_identity
        cols = schema.supporting_columns.meta_data
        instance_column = schema.instance_join.instance_column

        meta.require(
            instance_column,
            ident.primary_key_column,
            ident.wellplate_id_column,
            ident.assay_run_id_column,
            ident.well_number_column,
            cols.duration_seconds,
            cols.number_of_configurations,
            cols.sampling_frequency_hz,
            cols.gain,
            cols.lsb_uv,
            cols.hpf_hz,
            cols.start_time,
            cols.stop_time,
            cols.date,
        )

        analysis_types = self._analysis_types()
        found: dict[str, Recording] = {}

        for number, row in enumerate(meta.rows, start=2):
            where = f"{meta.name}!row{number}"
            folder_path = self.as_str(row[ident.primary_key_column], where=where)
            if folder_path is None:
                raise WorkbookError(
                    COLUMN_MISSING, f"{where}: {ident.primary_key_column!r} is missing"
                )

            instance = self.as_int(row[instance_column], where=f"{where}.{instance_column}")
            if instance is None:
                raise WorkbookError(COLUMN_MISSING, f"{where}: {instance_column!r} is missing")

            identity = RecordingIdentity(
                folder_path=folder_path,
                wellplate_id=self.as_str(row[ident.wellplate_id_column], where=where) or "",
                assay_run_id=self.as_str(row[ident.assay_run_id_column], where=where) or "",
                well_number=self.as_int(row[ident.well_number_column], where=where) or 0,
            )
            acquisition = Acquisition(
                duration_seconds=self.as_float(row[cols.duration_seconds], where=where),
                number_of_configurations=self.as_int(
                    row[cols.number_of_configurations], where=where
                ),
                sampling_frequency_hz=self.as_float(row[cols.sampling_frequency_hz], where=where),
                gain=self.as_float(row[cols.gain], where=where),
                lsb_uv=self.as_float(row[cols.lsb_uv], where=where),
                hpf_hz=self.as_float(row[cols.hpf_hz], where=where),
            )
            start = self.as_datetime(row[cols.date], row[cols.start_time], where=where)
            stop = self.as_datetime(row[cols.date], row[cols.stop_time], where=where)

            analysis_type = analysis_types.get(instance)
            existing = found.get(folder_path)
            if existing is None:
                instances = {analysis_type: (instance,)} if analysis_type else {}
                found[folder_path] = Recording(
                    identity=identity,
                    acquisition=acquisition,
                    start=start,
                    stop=stop,
                    instances=instances,
                )
            else:
                if analysis_type:
                    existing.instances[analysis_type] = (
                        existing.instances.get(analysis_type, ()) + (instance,)
                    )

        return found

    def _analysis_types(self) -> dict[int, str]:
        """Map Instance -> Analysis Type. The only sanctioned way to classify."""
        schema = self.schema
        sheet = self.sheet("analysis_parameters")
        instance_column = schema.instance_join.instance_column
        type_column = schema.instance_join.analysis_type_column
        sheet.require(instance_column, type_column)

        mapping: dict[int, str] = {}
        for number, row in enumerate(sheet.rows, start=2):
            where = f"{sheet.name}!row{number}"
            instance = self.as_int(row[instance_column], where=where)
            analysis_type = self.as_str(row[type_column], where=where)
            if instance is None or analysis_type is None:
                continue
            mapping[instance] = analysis_type
        return mapping

    def analysis_parameter_row(self, instance: int) -> dict[str, Any] | None:
        sheet = self.sheet("analysis_parameters")
        instance_column = self.schema.instance_join.instance_column
        for number, row in enumerate(sheet.rows, start=2):
            if self.as_int(row[instance_column], where=f"{sheet.name}!row{number}") == instance:
                return row
        return None

    def activity_row(self, instance: int) -> dict[str, Any] | None:
        sheet = self.sheet("activity_well_level")
        instance_column = self.schema.instance_join.instance_column
        sheet.require(instance_column)
        for number, row in enumerate(sheet.rows, start=2):
            if self.as_int(row[instance_column], where=f"{sheet.name}!row{number}") == instance:
                return row
        return None

    def __iter__(self) -> Iterator[Recording]:
        return iter(self.recordings().values())
